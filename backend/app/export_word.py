from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt


def generate(resume_json):
    """Returns a doc of the given resume JSON file."""
    # Prepare output file stream
    file_stream = BytesIO()

    # Create document and save it to files stream
    document = Document()
    set_layout(document)
    set_defaults(document)
    add_header(document, resume_json['header'])
    add_summary(document, resume_json['summary'])
    add_sections(document, resume_json['sections'])

    # Save document to files stream
    document.save(file_stream)
    file_stream.seek(0)

    return file_stream


def set_defaults(document):
    """Sets a default font and size for the document."""
    style = document.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = WD_LINE_SPACING.SINGLE


def set_layout(document):
    """Sets the page layout of a given Document object."""
    for section in document.sections:
        # Narrow margins
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

        # A4 size (8.27 x 11.69 inches)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)


def add_header(document, header):
    """Adds the header to the given resume document."""
    # Get header parameters
    name = header['name']
    email = header['email']
    mobile = header['mobile']
    portfolio = header['portfolio']
    linkedin = header['linkedin']

    # Name paragraph
    paragraph = document.add_paragraph()
    run = paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(16)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Email and mobile paragraph
    paragraph = document.add_paragraph(f'Email: {email}, Mobile: {mobile}')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Portfolio and LinkedIn paragraph
    paragraph = document.add_paragraph(f'Portfolio: {portfolio}, LinkedIn: {linkedin}')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Newline
    add_newline(document)


def add_summary(document, summary):
    # Add section title
    title = 'Summary'.upper()
    paragraph = document.add_paragraph()
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(12)

    # Add section contents
    paragraph = document.add_paragraph(summary)

    # Add newline
    add_newline(document)


def add_sections(document, sections):
    """Adds sections to the given resume document."""
    for index, section in enumerate(sections):
        add_section(document, section)

        # Add a newline if the section is not the last
        if index != len(sections) - 1:
            add_newline(document)


def add_section(document, section):
    """Add a particular section to the given resume document."""
    # Add section title
    title = section['title'].upper()
    vectorize = section['vectorize']
    checked = section['checked']
    entries = section['entries']

    if not checked:
        return None

    paragraph = document.add_paragraph()
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(12)

    # Add section contents
    checked_entries = [entry for entry in entries if entry['checked']]
    if vectorize == 'entries':
        if all(isinstance(entry['content'], str) for entry in entries):
            document.add_paragraph(', '.join(entry['content'] for entry in checked_entries))
        else:
            [
                add_entry(document, vectorize, entry['content'], index == len(entries) - 1)
                for index, entry in enumerate(checked_entries)
            ]

    if vectorize == 'bulletlist':
        [
            add_entry(document, vectorize, entry, index == len(entries) - 1)
            for index, entry in enumerate(checked_entries)
        ]


def add_entry(document, vectorize, entry_content, is_last):
    """Add an entry of a section to the given resume document."""
    # Get entry parameters
    name = entry_content['name']
    dates = entry_content['dates']

    # Add certification
    if 'institution' in entry_content and 'bulletlist' not in entry_content:
        institution = entry_content['institution']
        add_certification(document, name, dates, institution)

    # Add entries with headers, (probably subheaders?) and bulletpoints
    else:
        add_full_entry(
            document,
            name,
            dates,
            is_last,
            company=entry_content.get('company', None),
            bulletlist=entry_content.get('bulletlist', None),
        )


def add_certification(document, name, dates, institution):
    """Add a given certification of a section to the given resume document."""
    paragraph = document.add_paragraph('', style='List Bullet')
    run = paragraph.add_run(name)
    run.bold = True
    paragraph.add_run(f', {institution} - {dates}')


def add_full_entry(document, name, dates, is_last, company=None, bulletlist=None):
    """Add a full entry with headers (possibly subheaders?) and bulletpoints."""
    if bulletlist is None:
        bulletlist = []

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(7.27),  # right margin position
        alignment=WD_TAB_ALIGNMENT.RIGHT,
    )
    run = paragraph.add_run(f'{name}\t{dates}')
    run.bold = True

    # Company paragraph (if present)
    if company is not None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(company)
        run.bold = True

    # Bulletpoints
    for bulletpoint in bulletlist:
        if not bulletpoint['checked']:
            continue

        if 'string' in bulletpoint:
            bulletpoint_content = bulletpoint['string']

        if 'content' in bulletpoint:
            bulletpoint_content = bulletpoint['content']

        document.add_paragraph(bulletpoint_content, style='List Bullet')

    if not is_last:
        add_newline(document)


def add_newline(document):
    """Add a newline to the given resume document."""
    document.add_paragraph()
